#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Імпорт каталогу дакімакур з вигрузки товарів Prom (XLSX).

  python3 tools/import-prom.py export-products.xlsx

Що робить:
  1. бере з вигрузки всі дакімакури й згортає розміри одного малюнка в один дизайн;
  2. прибирає 18+, товари без подушки (лише наволочка) та id з tools/dakimakury-exclude.txt;
  3. розкладає дизайни по розділах, колекціях і персонажах (довідники в tools/prom_tables.py);
  4. зберігає назви, продажі й оригінали принтів для дизайнів, які вже є на сайті (dakimakury/data.js);
  5. пише tools/dakimakury-src.json. Сторінки з нього збирає tools/build-dakimakury.js.
"""
import sys, os, re, json, collections, subprocess
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import prom_tables as T
C = collections.Counter

# ---------- читання вигрузки ----------
def read_export(path):
    cache = path + '.json'
    if os.path.exists(cache) and os.path.getmtime(cache) >= os.path.getmtime(path):
        d = json.load(open(cache)); return d['prods'], d['groups']
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True); ps, gs = wb.worksheets[:2]
    rows = list(ps.iter_rows(values_only=True)); H = rows[0]; prods = []
    for r in rows[1:]:
        d = {H[i]: r[i] for i in range(51) if r[i] not in (None, '')}
        d['ch'] = {str(r[i]): ('' if r[i + 2] is None else str(r[i + 2])) for i in range(51, len(r) - 2, 3) if r[i]}
        prods.append(d)
    grows = list(gs.iter_rows(values_only=True)); GH = grows[0]
    groups = [{GH[i]: r[i] for i in range(6) if r[i] not in (None, '')} for r in grows[1:]]
    json.dump({'prods': prods, 'groups': groups}, open(cache, 'w'), ensure_ascii=False)
    return prods, groups

nm = lambda p: str(p.get('Назва_позиції_укр') or p.get('Назва_позиції') or '')
nmru = lambda p: str(p.get('Назва_позиції') or '')
imgs = lambda p: [u.strip() for u in str(p.get('Посилання_зображення', '')).split(',') if u.strip()]
pid = lambda u: int(u.split('/')[-1].split('_')[0])
size = lambda p: str(p['ch'].get('Размер подушки') or p['ch'].get('Розмір подушки') or '')
full = lambda p: 'наволочка' not in size(p).lower()
pkey = lambda p: str(p['Унікальний_ідентифікатор'])
NEW_PHOTO = 6e9          # фото з номером від цього це нові мокапи

SZ = r'(?:\d{2,3}\s*[*xх×]\s*\d{2,3})'
BOIL = [r'т[іо]льк[ио] наволочка', r'\bнаволочка\b',
        r'ан[иі]м[еєіяи]+\s+да[кн][иі]макура', r'да[кн][иі]макура', r'^\s*ан[иі]м[єе]\s+(?=[А-ЯІЇЄA-Z])', r'\(?\s*подушка[\s-]*об[нийі]+машка\s*\)?(\s+ростова)?',
        r'ростова[я]?\s+подушка[\s-]*об[нийі]+машка', r'подушка\s+ростова[я]?', r'ростова[я]?',
        r'\(\s*подушка\s*\)', r'\bподушка\b', r'\(\s*(габардин|велюр)\s*\)', r'\b(габардин|велюр)\b', r'для сн[уа]',
        r'\b(100|120|150|180)\s*,\s*' + SZ + r'\s*,\s*(33|40|50|60)\b', SZ + r'\s*,\s*(33|40|50|60)\b', SZ,
        r'\b(100|120|150|180)\s*,[\s,]*(33|40|50|60)\b', r'\bсм\b', r'\.0\b', r'\*']
def tidy(s):
    s = re.sub(r'\(\s*\)', ' ', s); s = re.sub(r'\s+,', ',', s); s = re.sub(r',(\s*,)+', ',', s)
    s = re.sub(r'\s+([)])', r'\1', s); s = re.sub(r'([(])\s+', r'\1', s)
    return re.sub(r'\s+', ' ', s).strip(' ,.-/–—:;')
def clean(s):
    s = ' ' + s + ' '
    for b in BOIL: s = re.sub(b, ' ', s, flags=re.I)
    return tidy(s)
core = lambda s: re.sub(r'[^0-9a-zа-яіїєґё]', '', clean(s).lower())

# ---------- рядки вигрузки -> дизайни ----------
def cluster(rows):
    parent = list(range(len(rows)))
    def find(x):
        while parent[x] != x: parent[x] = parent[parent[x]]; x = parent[x]
        return x
    cores = [core(nm(p)) for p in rows]; ph = [[pid(u) for u in imgs(p)] for p in rows]
    seen = collections.defaultdict(set)
    for i, p in enumerate(rows):
        for x in ph[i]: seen[x].add((p.get('Номер_групи'), cores[i]))
    generic = {x for x, v in seen.items() if len(v) >= 6}      # тканина, блискавка: одне фото в багатьох товарах
    first = {}
    for i, p in enumerate(rows):
        g = p.get('Номер_групи'); v = p.get('ID_групи_різновидів'); keys = []
        if v: keys.append(('v', g, v, cores[i]))
        keys += [('p', g, x, cores[i]) for x in ph[i] if x not in generic]
        for k in keys:
            if k in first:
                a, b = find(first[k]), find(i)
                if a != b: parent[b] = a
            else: first[k] = i
    out = collections.defaultdict(list)
    for i, p in enumerate(rows): out[find(i)].append(p)
    return list(out.values()), generic

# ---------- транслітерація для адрес ----------
TR = {'а':'a','б':'b','в':'v','г':'h','ґ':'g','д':'d','е':'e','є':'ye','ж':'zh','з':'z','и':'y','і':'i','ї':'yi','й':'y','к':'k','л':'l','м':'m',
      'н':'n','о':'o','п':'p','р':'r','с':'s','т':'t','у':'u','ф':'f','х':'kh','ц':'ts','ч':'ch','ш':'sh','щ':'shch','ь':'','ю':'yu','я':'ya',
      'ё':'yo','ъ':'','ы':'y','э':'e','’':'',"'":'','ʼ':'','`':''}
def slug(s):
    s = ''.join(TR.get(ch, ch) for ch in s.lower())
    s = s.replace('&', ' i ').replace('+', ' plus ')
    import unicodedata
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-')
def skey(s):
    """ключ для склеювання різних написань одного імені"""
    s = s.lower().replace('ё', 'е')
    for a, b in (('ї', 'и'), ('і', 'и'), ('й', 'и'), ('є', 'е'), ('э', 'е'), ('ґ', 'г'), ('ы', 'и'), ('ъ', ''), ('ь', '')): s = s.replace(a, b)
    s = re.sub(r'[^0-9a-zа-я]', '', s)
    return re.sub(r'(.)\1+', r'\1', s)

CATCHALL = {'anime2', 'games2', 'movies2', 'misc', 'kpop'}
SEC_CATCH = {'anime': 'anime2', 'games': 'games2', 'movies': 'movies2', 'kpop': 'kpop', 'other': 'misc', 'stars': 'celebs', 'sport': 'football', 'drinks': 'softdrinks'}
CATCH_H = {'anime2': 'Дакімакури: інше аніме', 'games2': 'Дакімакури з інших ігор', 'movies2': 'Дакімакури з інших фільмів і серіалів',
           'kpop': 'Дакімакури K-pop: інші гурти'}
SINGLE = {'ronaldo', 'gosling', 'revo'}                       # колекція про одну людину чи бренд, персонажів немає
SUB_LABEL = {'beer': 'Бренд', 'alcohol': 'Бренд', 'softdrinks': 'Бренд', 'anime2': 'Тайтл', 'games2': 'Тайтл', 'movies2': 'Тайтл', 'misc': 'Тайтл',
             'kpop': 'Гурт', 'celebs': 'Зірка', 'music': 'Зірка', 'ukr': 'Зірка', 'football': 'Гравець', 'f1': 'Гонщик', 'avto': 'Марка'}
MERGE_SUBS = {'genshin': {'Шеньхе': 'Шень Хе', 'Ганью': 'Гань Юй', 'Еола': 'Еула', 'Єлан': 'Є Лань'}}
# виправлення за id товару: куди насправді належить дизайн
FIX = {'2487336903': ('title', 'Honkai Impact 3rd'), '2688302281': ('col', 'twilight'), '2806433485': ('col', 'vocaloid')}
DRINK = (r'^(пиво|горілка|настоянка|віскі|вино|бальзам|наливка|лікер|коньяк|ром|джин|текіла|бренді|шампанське|сидр|вермут|абсент|самогон|'
         r'соджу|тонік|напій(\s+(енергетичний|слабоалкогольний|алкогольний|соковмісний|безалкогольний|газований|сильногазований))*|чай холодний|'
         r'сильногазований напій|енергетик|енергетичний напій|вода(\s+мінеральна)?|сік|квас|лимонад|кола)(\s+(ігристе|біле|червоне|рожеве|сухе))*\b[,\s]*')
ALCO = r'^(пиво|горілка|настоянка|віскі|вино|бальзам|наливка|лікер|коньяк|ром|джин|текіла|бренді|шампанське|сидр|вермут|абсент|самогон|соджу|напій (слабо)?алкогольний)\b'
BRAND2 = {'jack', 'captain', 'johnnie', 'coca', 'dr', 'mountain', 'red', 'нове', 'лавка', 'monster', 'ever', 'chupa', 'geo', 'перша', 'jim', 'old',
          'alfa', 'aston', 'land', 'mercedes', 'rolls', 'great', 'de', 'non', 'le', 'san', 's&r\'s', 'pisnya', 'proshyan', 'four', 'white', 'black',
          'козацька', 'легка', 'ukrainian', 'золоте', 'хлібний', 'grey', 'mogutni', 'royal', 'artisan', 'club', 'drink', 'лавка', 'stará', 'nemiroff'}
CAPS = {'Bmw': 'BMW', 'Gmc': 'GMC', 'Mg': 'MG', 'Byd': 'BYD', 'Vw': 'VW', 'Suv': 'SUV', 'Gt': 'GT', 'Gtr': 'GTR', 'Rs': 'RS', 'Amg': 'AMG', 'Kia': 'Kia',
        'Daf': 'DAF', 'Man': 'MAN', 'Uaz': 'УАЗ', 'Zaz': 'ЗАЗ', 'Vaz': 'ВАЗ', 'Gaz': 'ГАЗ', 'Ds': 'DS', 'Mini': 'MINI', 'Lwb': 'LWB'}
SHIP = {'hms', 'uss', 'ijn', 'kms', 'mnf', 'sn', 'rn', 'ffnf', 'roc', 'pran'}
LAT = re.compile(r'[a-z]', re.I); CYR = re.compile(r'[а-яіїєґё]', re.I)

def main():
    if len(sys.argv) < 2: sys.exit(__doc__)
    prods, groups = read_export(sys.argv[1])
    G = {str(g['Номер_групи']): g for g in groups}
    def gname(n, ru=False):
        g = G.get(str(n), {}); return (g.get('Назва_групи') if ru else (g.get('Назва_групи_укр') or g.get('Назва_групи'))) or ''
    def gpath(n):
        out = []; n = str(n); seen = set()
        while n in G and n not in seen:
            seen.add(n); out.append(gname(n)); n = str(G[n].get('Номер_батьківської_групи', ''))
        return list(reversed(out))
    dk = [p for p in prods if re.search(r'дак[иі]макур', (nmru(p) + ' ' + nm(p)).lower()) and size(p) != '20*60']
    print('рядків дакімакур у вигрузці:', len(dk))
    L = json.load(open(os.path.join(HERE, 'dakimakury-legacy.json')))
    LI = L['items']; LC = {c['k']: c for c in L['cols']}; SECK = [s['k'] for s in L['secs']]
    skip = set()
    for line in open(os.path.join(HERE, 'dakimakury-exclude.txt')):
        m = re.match(r'\s*(\d+)', line)
        if m: skip.add(m.group(1))
    exname = re.compile(T.EXCLUDE_NAME, re.I)
    # попередня збірка: щоб колекції не зникали й не зʼявлялися через кілька прибраних дизайнів, а старі адреси перенаправлялися
    prev_path = os.path.join(HERE, 'dakimakury-src.json')
    PREV = json.load(open(prev_path)) if os.path.exists(prev_path) else {'cols': [], 'redirects': {}}
    prev_titles = {c['t'] for c in PREV['cols'] if c['k'] not in LC}

    designs, generic = cluster(dk)
    designs = [v for v in designs if any(full(p) and p.get('Наявність') != '-' for p in v)]
    # той самий малюнок двома оголошеннями: окремо «лише 150×50» і окремо решта розмірів
    bucket = collections.defaultdict(list)
    for v in designs: bucket[(v[0].get('Номер_групи'), core(nm(v[0])))].append(v)
    drop = set()
    for b in bucket.values():
        only50 = [v for v in b if {size(p) for p in v if full(p)} == {'50*150'}]
        rest = [v for v in b if '50*150' not in {size(p) for p in v if full(p)}]
        for v in only50[:len(rest)]:
            if not any(pkey(p) in LI for p in v): drop.add(id(v))
    designs = [v for v in designs if id(v) not in drop]
    print('унікальних дизайнів, що продаються:', len(designs))

    # усі назви тайтлів, за якими впізнаємо тайтл у товарах збірних груп
    title_alias = []
    for g, al in T.ALIASES.items():
        t = T.GROUP_TO_COL.get(g)
        for a in al:
            if len(a) >= 5: title_alias.append((a.lower(), ('col', t) if t else ('title', T.RENAME.get(g, g))))
    for k, al in T.ROUTE.items(): title_alias += [(a.lower(), ('col', k)) for a in al]
    title_alias += [(a, ('title', t) if t else ('skip', '')) for a, t in T.TITLE_ROUTE]
    title_alias.sort(key=lambda x: -len(x[0]))
    def sec_of(title):
        return ('games' if title in T.GAMES else 'movies' if title in T.MOVIES else 'sport' if title in T.SPORT else 'other' if title in T.OTHER else 'anime')

    W = r'0-9a-zа-яіїєґё'
    def has(a, text): return re.search(r'(?<![' + W + r'])' + re.escape(a) + r'(?![' + W + r'])', text) is not None
    stats = C(); recs = []
    for v in designs:
        fr = [p for p in v if full(p)] or v
        gid = C(p.get('Номер_групи') for p in fr).most_common(1)[0][0]
        path = gpath(gid); g = path[-1] if path else ''
        ids = [pkey(p) for p in v]
        leg = sorted([i for i in ids if i in LI and i not in skip], key=lambda i: -LI[i]['sold'])
        if skip & set(ids) and not leg and (any(i in LI for i in ids) or next((pkey(p) for p in v if size(p) == '40*120' and p.get('Наявність') != '-'), pkey(fr[0])) in skip):
            stats['викинуто: список винятків'] += 1; continue
        names = ' | '.join(sorted({nm(p) + ' ' + nmru(p) for p in fr}))
        if set(path) & T.EXCLUDE_GROUPS: stats['викинуто: групи 18+ і дитячі тайтли'] += 1; continue
        # підрозділ «Еротичні подарунки» продавець ставить не лише відвертим малюнкам, тому дизайни, які вже є на сайті, через нього не прибираємо
        if not leg and any('Eroticheskie' in str(p.get('Посилання_підрозділу')) for p in v): stats['викинуто: еротичний підрозділ Prom'] += 1; continue
        if exname.search(names): stats['викинуто: за словами в назві'] += 1; continue
        # фото: найновіший мокап серед варіантів
        def rank(p):
            im = imgs(p); real = [u for u in im if pid(u) not in generic]
            return (bool(real) and pid(real[0]) >= NEW_PHOTO, bool(leg) and pkey(p) == leg[0], size(p) == '40*120', len(real))
        best = max(fr, key=rank)
        im = imgs(best); real = [u for u in im if pid(u) not in generic]
        if not real: stats['викинуто: немає фото малюнка'] += 1; continue
        pics = []
        for u in real + [u for u in im if pid(u) in generic]:
            u = u.split('/')[-1]; u = u[:-4] if u.endswith('.jpg') else 'https://images.prom.ua/' + u
            if u not in pics: pics.append(u)
        base = clean(nm(best)); low = (base + ' ' + clean(nmru(best))).lower()

        # ---- куди ----
        where = None
        for i in ids:
            if i in FIX: where = FIX[i]
        if not where and leg and LI[leg[0]]['c'] not in CATCHALL: where = ('col', LI[leg[0]]['c'])
        if not where:
            if 'K-POP' in path or g == 'K-pop':
                k = T.GROUP_TO_COL.get(g)
                if not k and re.search(r'stray\s*kids|стрей', low): k = 'straykids'
                if not k and re.search(r'\bbts\b|\bбтс\b', low): k = 'bts'
                if k: where = ('col', k)
                else:
                    band = T.BAND_ALIAS.get(g, g) if g != 'K-pop' else None
                    if not band:
                        for b in T.BANDS:
                            if b.lower() in low: band = T.BAND_ALIAS.get(b, b); break
                    if not band:
                        for person, b in T.PERSON_BAND.items():
                            if re.search(r'(^|[\s,(])' + re.escape(person) + r'($|[\s,)])', low): band = b; break
                    where = ('title', band) if band else ('col', 'kpop')
            elif g in T.GROUP_TO_COL:
                k = T.GROUP_TO_COL[g]
                if k == 'alcohol' and re.search(r'\bпиво\b', low): k = 'beer'
                if k == 'softdrinks' and 'revo' in low: k = 'revo'
                if k == 'football' and re.search(r'роналд|ronaldo', low): k = 'ronaldo'
                if k == 'celebs':
                    if 'гослінг' in low: k = 'gosling'
                    for rx, w in T.CELEB_ROUTE:
                        if re.search(rx, low): where = w; break
                where = where or ('col', k)
            elif g == 'Інші' or not g:
                for a, w in title_alias:
                    if has(a, low): where = w; break
                if where and where[0] == 'skip': stats['викинуто: групи 18+ і дитячі тайтли'] += 1; continue
                if not where:
                    b = base.lower()
                    if re.search(DRINK, b): where = ('col', 'beer' if re.search(r'^пиво', b) else 'alcohol' if re.search(ALCO, b) else 'softdrinks')
                    elif re.search(r'дак[иі]макура,\s*подушка ростова', nm(best), re.I): where = ('col', 'celebs')
                    elif re.search(r'фуррі|furry', b): where = ('col', 'misc')
                    else: where = ('col', 'anime2')
            else: where = ('title', T.RENAME.get(g, g))
        recs.append(dict(v=v, best=best, g=g, path=path, ids=ids, leg=leg, pics=pics[:4], base=base, where=where, gid=gid))
    for k, n in stats.most_common(): print(' ', k + ':', n)

    # ---- тайтли: великі стають колекціями, малі йдуть у збірну колекцію розділу ----
    tcount = C(r['where'][1] for r in recs if r['where'][0] == 'title')
    cols = [dict(c) for c in L['cols']]; byk = {c['k']: c for c in cols}
    for c in cols: c['s'] = SECK[c['s']]; c.pop('subs', None); c.pop('sp', None); c['h'] = CATCH_H.get(c['k'], c['h'])
    used = {c['p'] for c in cols} | {s['p'] for s in L['secs']}
    tkey = {}
    for t, n in tcount.most_common():
        sec = 'kpop' if t in set(T.BANDS) | set(T.BAND_ALIAS.values()) else sec_of(t)
        if n >= 10 or (t in prev_titles and n >= 5):      # уже опублікована колекція лишається, доки в ній є хоча б 5 дизайнів
            p = slug(t); k = p
            while p in used: p += '-dakimakury'
            used.add(p); tkey[t] = k
            c = dict(k=k, t=t, h=T.H1.get(t, 'Дакімакури ' + t), s=sec, top=[], p=p); cols.append(c); byk[k] = c
        else: tkey[t] = None
    for r in recs:
        kind, x = r['where']
        if kind == 'col': r['c'] = x; r['title'] = ''
        elif tkey[x]: r['c'] = tkey[x]; r['title'] = ''
        else:
            sec = 'kpop' if x in set(T.BANDS) | set(T.BAND_ALIAS.values()) else sec_of(x)
            r['c'] = SEC_CATCH[sec]; r['title'] = x
        if r['c'] not in byk: raise SystemExit('невідома колекція ' + r['c'])
    print('колекцій:', len(cols), '(нових %d)' % (len(cols) - len(L['cols'])))

    # ---- назви ----
    def aliases_for(r):
        g = r['g']; al = list(T.ALIASES.get(g, []))
        al += [g, T.RENAME.get(g, ''), gname(r['gid'], True), byk[r['c']]['t'], r['title']]
        al += [x.strip() for x in re.split(r'[()/:]', g) if len(x.strip()) >= 4]
        for p in r['v']:
            n = str(p.get('Особисті_нотатки') or '')
            if len(n) >= 4 and not re.search(r'dak|коп', n, re.I): al.append(n)
        return sorted({a for a in al if a and len(a) >= 2}, key=lambda a: -len(a))
    def derive(r):
        n = ' ' + r['base'] + ' '
        emptied = False
        if r['c'] not in ('beer', 'alcohol', 'softdrinks', 'revo', 'celebs', 'music', 'ukr', 'gosling', 'ronaldo', 'memes'):
            for a in aliases_for(r):
                n2 = re.sub(r'(?<![0-9a-zа-яіїєґё])[(\[]?\s*' + re.escape(a) + r'\s*[)\]]?(?![0-9a-zа-яіїєґё])', ' ', n, flags=re.I)
                if tidy(re.sub(r'№?\s*\d+', '', n2)): n = n2
                elif n2 != n: emptied = True
        r['nochar'] = emptied
        n = tidy(n)
        n = re.sub(r'\s*[(\[]\s*[A-ZА-ЯІЇЄ]\s*[)\]]$', '', n)                 # (E), (D): варіант малюнка
        if byk[r['c']]['s'] in ('stars', 'kpop', 'sport', 'movies') and CYR.search(n.split('(')[0].split(',')[0]):
            n = re.sub(r'\s*\(\s*[^()а-яіїєґё]*[a-zA-Z][^()а-яіїєґё]*\)', '', n)        # «Бред Пітт (Brad Pitt)»
            n = re.sub(r',\s*[A-Za-z][^,а-яіїєґё]*$', '', n)                          # «Біллі Айліш, Billie Eilish»
        if n.count(')') > n.count('('): n = n.replace(')', '', 1)
        w = n.split()
        if len(w) >= 2 and len(w) % 2 == 0 and w[:len(w) // 2] == w[len(w) // 2:]: n = ' '.join(w[:len(w) // 2])
        n = re.sub(r'(?<=\D)[\s,]+\d{1,3}$', '', n) if len(n.split()) > 1 else n   # номер малюнка в кінці
        n = re.sub(r'^([А-ЯІЇЄа-яіїє]+),\s+', r'\1 ', n) if re.search(DRINK, n.lower()) else n
        n = tidy(re.sub(r'\s+-\s+', ' ', re.sub(r'^[\s\-–,]+|[\s\-–,]+$', '', n)))
        if r['c'] == 'avto' or r['title'] == 'Авто': n = ' '.join(CAPS.get(w, w) for w in n.split())
        n = re.sub(r'(?<=[а-яіїєґ])(\d)$', r' №\1', n)
        if not n: r['nochar'] = True
        n = n or r['title'] or byk[r['c']]['t']
        return n[0].upper() + n[1:]
    SZL = re.compile(r'\s*\b(100|120|150|180),\s*(33|40|50|60)\b')
    for r in recs:
        if r['leg']:
            li = LI[r['leg'][0]]
            r['name'] = tidy(SZL.sub('', li['name'])); r['ru'] = tidy(SZL.sub('', li['ru']))
            m = re.match(r'^(.+?) \1( №\d+)?$', r['name'])                  # «Хатсуне Міку Хатсуне Міку №3»
            if m: r['name'] = m.group(1) + (m.group(2) or '')
            r['moved'] = li['c'] in CATCHALL and r['c'] not in CATCHALL
            if r['moved']:                       # раніше лежав у збірній колекції: назву тайтлу з імені прибираємо
                r['base'] = re.sub(r'\s*№\s*\d+$', '', r['name']); r['name'] = derive(r)
            r['sold'] = sum(LI[i]['sold'] for i in r['leg']); r['id'] = r['leg'][0]; r['nochar'] = False
            o = next((LI[i] for i in r['leg'] if LI[i]['orig']), None)
            r['orig'] = o['orig'] if o else []; r['og'] = o['og'] if o else 0
        else:
            r['name'] = derive(r); r['sold'] = 0; r['orig'] = []; r['og'] = 0
            b = r['best']; r['id'] = pkey(next((p for p in r['v'] if size(p) == '40*120' and p.get('Наявність') != '-'), b))
            have = set(re.findall(r'[0-9a-zа-яіїєґё]+', (r['name'] + ' ' + byk[r['c']]['t']).lower()))
            ru = [w for w in re.findall(r"[0-9a-zа-яіїєґё'’:!.\-]+", clean(nmru(b)).lower()) if w.strip("'’:!.-") not in have]
            ru = [w for w in ru if len(w) > 1 and not re.match(r'^[\d.,]+(л|мл)?$', w)]
            if byk[r['c']]['s'] in ('drinks', 'other'): ru = ru[:2]
            extra = [r['title'].lower()] if r['title'] else []
            r['ru'] = ' '.join(dict.fromkeys(ru + extra))[:90].strip()

    # ---- персонажі, бренди, тайтли всередині колекції ----
    def char(name, c):
        n = re.sub(r'\s*№\s*\d+', '', name)
        if c in ('beer', 'alcohol', 'softdrinks'):
            n = re.sub(DRINK, '', n, flags=re.I); m = re.match(r'[«"“]([^»"”]+)[»"”]', n)
            if m: return m.group(1).strip()
            w = [x.strip('«»"“”,.') for x in n.split()]
            if not w or not w[0] or re.match(r'^[\d.,]+', w[0]): return ''
            return ' '.join(w[:2]) if w[0].lower() in BRAND2 and len(w) > 1 else w[0]
        if c == 'avto':
            w = n.split(); return (' '.join(w[:2]) if w and w[0].lower() in BRAND2 and len(w) > 1 else (w[0] if w else ''))
        n = re.sub(r'\([^()]*\)', ' ', n); n = re.split(r'\s[-–—]\s|,|\s/\s|/|\sта\s|\sі\s|\sи\s|\s&\s|\sx\s|\sх\s', n)[0]
        n = re.sub(r'(?<=[а-яіїєґa-z])\d+$', '', n.strip()); n = re.sub(r'\s+\d{1,3}$', '', n)
        w = n.split()
        if len(w) > 1 and w[0].lower() in SHIP: w = w[1:]
        if len(w) > 1:                                              # «Фуріна Furina», «Райден Shogun»: лишаємо перше написання
            first = 'c' if CYR.search(w[0]) else 'l' if LAT.search(w[0]) else ''
            keep = []
            for x in w:
                sc = 'c' if CYR.search(x) else 'l' if LAT.search(x) else first
                if sc != first: break
                keep.append(x)
            w = keep or w
        return tidy(' '.join(w))
    bycol = collections.defaultdict(list)
    for r in recs: bycol[r['c']].append(r)
    redirects = {}; gone = []
    for c in cols:
        k = c['k']; rs = bycol.get(k, []); lc = LC.get(k, {})
        lsubs = list(lc.get('subs') or []); lsp = dict(zip(lsubs, lc.get('sp') or []))
        merge = MERGE_SUBS.get(k, {})
        istitle = k in CATCHALL
        if istitle and k != 'kpop': lsubs = [s for s in lsubs if s not in ('Інші', 'Авто')]
        if k == 'kpop': lsubs = []
        keys = [(skey(s), s) for s in sorted(lsubs, key=lambda s: -len(s)) if s not in merge]
        for r in rs:
            s = ''
            if k in SINGLE: pass
            elif istitle: s = r['title'] or (LI[r['leg'][0]]['label'] if r['leg'] and k != 'kpop' else '')
            else:
                if r['leg'] and not r.get('moved'): s = merge.get(LI[r['leg'][0]]['sub'], LI[r['leg'][0]]['sub'])
                if not s:
                    nk = skey(r['name'])
                    s = next((orig for kk, orig in keys if len(kk) >= 4 and kk in nk), '') or char(r['name'], k)
            r['sub'] = s if len(skey(s)) >= 2 and not (r.get('nochar') and not istitle) and skey(s) != skey(c['t']) else ''
        # склеюємо різні написання
        grp = collections.defaultdict(C)
        for r in rs:
            if r['sub']: grp[skey(r['sub'])][r['sub']] += 1
        ks = sorted(grp, key=lambda x: -sum(grp[x].values())); canon = {}
        def near(a, b):
            if abs(len(a) - len(b)) > 1 or min(len(a), len(b)) < 6: return False
            if len(a) == len(b): return sum(x != y for x, y in zip(a, b)) == 1
            if len(a) < len(b): a, b = b, a
            return any(a[:i] + a[i + 1:] == b for i in range(len(a)))
        if len(ks) <= 1500 and not istitle and k not in ('avto', 'beer', 'alcohol', 'softdrinks'):
            for i, a in enumerate(ks):
                if a in canon: continue
                for b in ks[i + 1:]:
                    if b not in canon and near(a, b): canon[b] = a
        disp = {}; pc = next((x for x in PREV['cols'] if x['t'] == c['t']), None)
        pv = {skey(n): (n, sl) for n, sl in zip(pc.get('subs') or [], pc.get('sp') or [])} if pc else {}
        for a in ks:
            root = canon.get(a, a); grp[root].update(grp[a]) if root != a else None
        for a in ks:
            if a in canon: continue
            leg_name = next((s for s in lsubs if skey(merge.get(s, s)) == a), None)
            # написання й адреса з попередньої збірки лишаються, навіть якщо після чистки частіше трапляється інше написання
            disp[a] = merge.get(leg_name, leg_name) if leg_name else pv.get(a, (None,))[0] or grp[a].most_common(1)[0][0]
        cnt = C()
        for r in rs:
            if r['sub']:
                a = skey(r['sub']); a = canon.get(a, a); r['sub'] = disp[a]; cnt[r['sub']] += 1
        order = [merge.get(s, s) for s in lsubs if cnt.get(merge.get(s, s), 0) >= 2]
        order = list(dict.fromkeys(order)) + [s for s, n in cnt.most_common() if n >= 2 and s not in order and s not in {merge.get(x, x) for x in lsubs}]
        usedp = set(); sp = []
        for s in order:
            p = next((lsp[x] for x in lsubs if merge.get(x, x) == s and lsp.get(x)), '') or pv.get(skey(s), (None, ''))[1] or slug(s) or 'p'
            while p in usedp: p += '-2'
            usedp.add(p); sp.append(p)
        idx = {s: i for i, s in enumerate(order)}
        for r in rs: r['k'] = idx.get(r['sub'], -1)
        c['subs'] = order; c['sp'] = sp
        c['sl'] = SUB_LABEL.get(k, lc.get('sl') or 'Персонаж')
        for s, p in lsp.items():
            if p and merge.get(s, s) not in idx: gone.append((k, s, '%s/%s' % (lc['p'], p)))
            elif p and p != sp[idx[merge.get(s, s)]]: redirects['%s/%s' % (lc['p'], p)] = '%s/%s' % (c['p'], sp[idx[merge.get(s, s)]])   # два написання злилися в одне

    # ---- однакові назви в колекції: додаємо номери ----
    for k, rs in bycol.items():
        byname = collections.defaultdict(list)
        for r in rs: byname[r['name']].append(r)
        taken = collections.defaultdict(set)
        for n in byname:
            m = re.match(r'^(.*?)\s*№\s*(\d+)$', n)
            if m: taken[m.group(1)].add(int(m.group(2)))
        for n, lst in byname.items():
            if len(lst) < 2 or re.search(r'№\s*\d+$', n): continue
            lst.sort(key=lambda r: (not r['leg'], -r['sold'], r['id'])); i = 0
            for r in lst:
                i += 1
                while i in taken[n]: i += 1
                taken[n].add(i); r['name'] = '%s №%d' % (n, i)

    # старі адреси підкатегорій, яких більше немає: ведемо туди, куди переїхала більшість їхніх дизайнів
    for k, sname, old in gone:
        to = C()
        for r in recs:
            if any(LI[i]['c'] == k and LI[i]['sub'] == sname for i in r['leg']): to[(r['c'], r['k'])] += 1
        if to:
            (ck, sk_), _n = to.most_common(1)[0]; c = byk[ck]
            redirects[old] = c['p'] + ('/' + c['sp'][sk_] if sk_ >= 0 and (k == 'kpop' or k not in CATCHALL) else '')
        else: redirects[old] = LC[k]['p']
    cols = [c for c in cols if bycol.get(c['k'])]
    # сторінки попередньої збірки, яких тепер немає
    now = {c['p'] for c in cols} | {'%s/%s' % (c['p'], x) for c in cols for x in c['sp'] if x}
    nowcol = {c['t']: c for c in cols}; secp = {s['k']: s['p'] for s in L['secs']}
    for old, to in (PREV.get('redirects') or {}).items(): redirects.setdefault(old, to)
    for pc in PREV['cols']:
        home = nowcol.get(pc['t'])
        if pc['p'] not in now:                       # колекція стала підкатегорією збірної або зникла
            to = secp.get(pc['s'], '')
            catch = byk.get(SEC_CATCH.get(pc['s'], ''))
            if catch and catch in cols:
                to = catch['p'] + ('/' + catch['sp'][catch['subs'].index(pc['t'])] if pc['t'] in catch['subs'] else '')
            redirects[pc['p']] = to
        for x in pc.get('sp') or []:
            if x and '%s/%s' % (pc['p'], x) not in now:
                redirects['%s/%s' % (pc['p'], x)] = home['p'] if home and home['p'] in now else redirects.get(pc['p'], secp.get(pc['s'], ''))
    redirects = {a: b for a, b in redirects.items() if a not in now}
    for a in list(redirects):                        # ланцюжки: стара адреса одразу веде на живу сторінку
        seen = set()
        while redirects[a] in redirects and redirects[a] not in seen: seen.add(redirects[a]); redirects[a] = redirects[redirects[a]]
    for c in L['cols']:
        if c['k'] not in {x['k'] for x in cols}: redirects[c['p']] = L['secs'][c['s']]['p']
    live = {c['k'] for c in cols}
    items = [[r['id'], r['c'], r['name'], r['ru'], r['pics'], r['sold'], r['title'] if r['c'] in CATCHALL else '', r['k'], r['orig'], r['og']] for r in recs if r['c'] in live]
    # спершу те, що купують; далі дизайни з іменем персонажа, безіменні наприкінці
    named = {r['id']: (0 if r['k'] >= 0 else 1, 1 if r.get('nochar') else 0) for r in recs}
    items.sort(key=lambda i: (-i[5],) + named[i[0]] + (i[0],))
    import datetime
    out = dict(v=datetime.date.today().isoformat(), secs=L['secs'], cols=cols, items=items, redirects=redirects)
    json.dump(out, open(os.path.join(HERE, 'dakimakury-src.json'), 'w'), ensure_ascii=False, separators=(',', ':'))
    print('дизайнів у каталозі:', len(items), '| з них уже були на сайті:', sum(1 for r in recs if r['leg']))
    print('підкатегорій зі сторінками:', sum(len(c['subs']) for c in cols), '| перенаправлень зі старих адрес:', len(redirects))
    lost = set(LI) - {i for r in recs for i in r['ids']}
    print('дизайнів сайту, яких немає серед імпортованих:', len(lost))
    json.dump(sorted(lost), open(os.path.join(HERE, '.import-lost.json'), 'w'))

if __name__ == '__main__':
    main()
